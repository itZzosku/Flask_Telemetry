"""Velocity import + statistics. Synthetic data only — the repo is public, real logs stay off it."""
import datetime as dt
import io
import math

import openpyxl
import pytest

import diary
from diary import chrono
from diary.chrono_import import (apply_weather, hms, num, parse_file, parse_kestrel, parse_load_name,
                                 parse_log_sheet, parse_shotview_sheet)
from diary.chrono_stats import stats, temp_sensitivity


# ------------------------------------------------------------------ helpers
def shotview_rows(speeds, title="Rifle Bullet 120,0 gr", unit="MPS", sd=None, date="OCTOBER 15, 2026 13:23"):
    avg = sum(speeds) / len(speeds)
    pop = math.sqrt(sum((v - avg) ** 2 for v in speeds) / len(speeds))
    fi = lambda x: f"{x:.1f}".replace(".", ",")
    rows = [[title] + [""] * 8,
            ["#", f"Speed ({unit})", f"Δ Avg ({unit})", "KE (J)", "Power Factor (N⋅s)", "Time", "Clean Bore",
             "Cold Bore", "Shot Notes"]]
    for i, v in enumerate(speeds):
        rows.append([str(i + 1), fi(v), fi(v - avg), "2\xa0720,8", "6,5", f"13:2{i}:00", "", "x" if i == 0 else "", ""])
    rows += [["-"], ["AVERAGE SPEED", fi(avg)], ["STD DEV", fi(pop if sd is None else sd)],
             ["SPREAD", fi(max(speeds) - min(speeds))], ["Projectile Weight (GRAINS)", "120,0"],
             ["Session Note", ""], ["-"], ["Date", date]]
    return rows


def log_workbook(sheets):
    """sheets = {name: dict(date, shots=[(t, v, temp, rh, hpa, charge)], shift=0, bullet=139)}"""
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    wb.create_sheet("TEMPLETE")
    for name, d in sheets.items():
        ws = wb.create_sheet(name)
        sh = d.get("shift", 0)
        ws["B2"], ws["C2"], ws["E2"] = "Time", "Date", "Average \nTemperature (C°)"
        ws.cell(2, 17 + sh, "Bullet weight\n (gr)")
        ws["B3"], ws["C3"] = d.get("start", "15:30"), d["date"]
        ws.cell(3, 17 + sh, d.get("bullet", 139.0))
        hdr = ["Shot Number", "Timestamp\nGarmin", "Speed \n(m/s)", "Temperature \n(C°)", "Humidity \n(% RH)",
               "Pressure \n( hPa)", "Timestamp\nKestrel"]
        for i, h in enumerate(hdr):
            ws.cell(5, 2 + i, h)
        if any(len(s) > 5 for s in d["shots"]):
            ws.cell(5, 20, "Charge weight \n(gr)")
        ws.cell(6, 10, "Min")                     # side table must be ignored
        for r, s in enumerate(d["shots"]):
            ws.cell(6 + r, 2, r + 1)
            ws.cell(6 + r, 3, s[0])
            ws.cell(6 + r, 4, s[1])
            ws.cell(6 + r, 5, s[2])
            ws.cell(6 + r, 6, s[3])
            ws.cell(6 + r, 7, s[4])
            if len(s) > 5 and s[5] is not None:
                ws.cell(6 + r, 20, s[5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ parsing basics
def test_num_and_time():
    assert num("768,1") == 768.1
    assert num("2\xa0720,8") == 2720.8
    assert num("--") is None and num("") is None and num(None) is None
    assert num(841.5) == 841.5
    assert hms("13.00") == "13:00:00"
    assert hms(dt.time(18, 0)) == "18:00:00"
    assert hms("15:33:32") == "15:33:32"
    assert hms("x") is None


def test_load_names():
    h = parse_load_name("T3X Super Varmint, 9 g Scenar, Lapua brass, 44.0 gr")
    assert (h["gun_text"], h["bullet"], h["brass"], h["charge_gr"]) == ("T3X Super Varmint", "9 g Scenar", "Lapua", 44.0)
    h = parse_load_name("M_39, 6,5 g HPCE, Lapua brass, 43.0 gr")         # decimal comma inside a part
    assert h["bullet"] == "6,5 g HPCE" and h["charge_gr"] == 43.0
    h = parse_load_name("M_39, 12 g Scenar, Lapua brass, test strings")
    assert h["charge_gr"] is None and h["charge_label"] == "test strings"
    h = parse_load_name("T1X, Geco .22lr Rifle 2,6g")
    assert h["kind"] == "factory" and h["factory"] == "Geco .22lr Rifle 2,6g" and round(h["bullet_gr"]) == 40
    h = parse_load_name("T1X, SK 40gr. Long Range Match")
    assert h["bullet_gr"] == 40.0
    h = parse_load_name("T3X_Super_Varmint_9_g_Scenar_Lapua_brass_44.0_gr")
    assert h["brass"] == "Lapua" and h["charge_gr"] == 44.0


# ------------------------------------------------------------------ ShotView
def test_shotview_sheet():
    s = parse_shotview_sheet("Rifle_Bullet_2026-10-15_13-42_2", shotview_rows([836.5, 836.6, 831.1]))
    assert s["date"] == "2026-10-15" and s["started_at"] == "13:23"
    assert s["bullet_gr"] == 120.0 and len(s["shots"]) == 3
    assert s["shots"][0]["cold_bore"] == 1 and s["shots"][0]["t"] == "13:20:00"
    assert s["check"]["ok"] is True and not s["warnings"]


def test_shotview_old_app_sd_zero_and_fps():
    s = parse_shotview_sheet("x", shotview_rows([1000.0, 1010.0, 1005.0], unit="FPS", sd=0.0))
    assert s["check"]["ok"] is True                       # SD 0,0 export bug is ignored
    assert s["unit_in"] == "fps" and abs(s["shots"][0]["v_ms"] - 304.8) < 0.01


def test_shotview_mismatch_flagged():
    rows = shotview_rows([800.0, 802.0, 804.0])
    rows[-7][1] = "790,0"                                 # wrong AVERAGE SPEED in the file
    s = parse_shotview_sheet("x", rows)
    assert s["check"]["ok"] is False and s["warnings"][0]["code"] == "check_mismatch"


# ------------------------------------------------------------------ own workbooks
def test_workbook_day_and_ladder_split():
    data = log_workbook({
        "13072024": {"date": dt.datetime(2024, 7, 13), "shots": [("15:33:32", 847.0, 23.0, 46.4, 1004.4),
                                                                   ("15:34:39", 842.6, 23.0, 46.4, 1004.4),
                                                                   ("15:34:48", 840.7, 23.0, 46.4, 1004.4)]},
        "14022026": {"date": dt.datetime(2026, 2, 14), "shots": [("18:41:52", 809.7, -12, 89, 1008, 31.0),
                                                                   ("18:42:08", 823.9, -12, 89, 1008, None),
                                                                   ("18:44:35", 857.7, -12, 89, 1008, 33.0),
                                                                   ("18:44:53", 856.6, -12, 89, 1008, None)]},
    })
    res = parse_file("T3X Super Varmint, 6.5 g Scenar, Lapua brass, x0.0 gr.xlsx", data,
                     "Velocity/N133/T3X Super Varmint/T3X Super Varmint, 6.5 g Scenar, Lapua brass, x0.0 gr.xlsx")
    days = {(s["date"], s["charge_gr"]): s for s in res["strings"]}
    assert set(days) == {("2024-07-13", None), ("2026-02-14", 31.0), ("2026-02-14", 33.0)}
    one = days[("2024-07-13", None)]
    assert one["bullet_gr"] == 139.0 and one["shots"][0]["hpa"] == 1004.4
    assert one["hints"]["powder"] == "N133" and "T3X Super Varmint" in one["hints"]["path"]
    assert len(days[("2026-02-14", 31.0)]["shots"]) == 2      # charge carried down


def test_workbook_shifted_columns_fixed_twin_and_swaps():
    shots = [("10:57:16", 736.4, 2, 84, 994), ("10:58:23", 691.1, 2, 84, 994), ("10:58:45", 682.1, 2, 84, 994)]
    data = log_workbook({
        "29122024": {"date": dt.datetime(2024, 12, 29), "shots": shots},
        "29122024 Fixed": {"date": dt.datetime(2024, 12, 29), "shots": shots[1:], "shift": 1},
        "14042025": {"date": dt.datetime(2025, 4, 14), "shots": [("19:00:00", 950.0, 8.4, 999.0, 87.0)] * 2},
    })
    res = parse_file("M_39, 12 g Scenar, Lapua brass, test strings.xlsx", data)
    sheets = {s["sheet"]: s for s in res["strings"]}
    assert "29122024" not in sheets and "29122024 Fixed" in sheets
    assert sheets["29122024 Fixed"]["bullet_gr"] == 139.0           # found by label, one column right
    w = sheets["14042025"]["warnings"]
    assert w and w[0]["code"] == "weather_range" and w[0]["swap"] is True
    assert any(n["code"] == "superseded" for n in res["notes"])


def test_year_typo_flag():
    mk = lambda d, t: {"date": d, "shots": [(t, 890.0, -4, 80, 1000), (t[:-1] + "9", 891.0, -4, 80, 1000)]}
    data = log_workbook({"05012024": mk(dt.datetime(2024, 1, 5), "12:24:50"),
                         "09012025": mk(dt.datetime(2025, 1, 9), "14:07:00"),
                         "22012025": mk(dt.datetime(2025, 1, 22), "17:46:00")})
    res = parse_file("M_39, 8 g FMJ, Lapua brass, 43.0 gr.xlsx", data)
    flagged = [s for s in res["strings"] if any(w["code"] == "date_outlier" for w in s["warnings"])]
    assert [s["date"] for s in flagged] == ["2024-01-05"]
    assert flagged[0]["warnings"][-1]["suggest"] == "2025-01-05"


# ------------------------------------------------------------------ Kestrel
KESTREL = ('"Device Name","Drop D3"\n"Serial Number","0000000"\n'
           '"FORMATTED DATE_TIME","Temperature","Relative Humidity","Station Pressure","Heat Index"\n'
           '"YYYY-MM-DD HH:MM:SS","°C","%","mb","°C"\n'
           '"2024-07-13 03:30:00 pm","23,4","42,9","1004,4","26,6"\n'
           '"2024-07-13 03:40:00 pm","23,0","39,2","1004,3","26,6"\n').encode()


def test_kestrel_merge_fills_only_empty():
    recs = parse_kestrel(KESTREL)
    assert recs[0] == {"date": "2024-07-13", "t": "15:30:00", "temp_c": 23.4, "rh": 42.9, "hpa": 1004.4}
    s = {"date": "2024-07-13", "shots": [
        {"t": "15:38:00", "temp_c": None, "rh": None, "hpa": None},
        {"t": "15:31:00", "temp_c": 20.0, "rh": None, "hpa": None},     # manual value stays
        {"t": "16:30:00", "temp_c": None, "rh": None, "hpa": None}]}    # outside ±10 min
    assert apply_weather([s], recs) == 2
    assert s["shots"][0]["rh"] == 39.2 and s["shots"][1]["temp_c"] == 20.0 and s["shots"][1]["rh"] == 42.9
    assert s["shots"][2]["temp_c"] is None


# ------------------------------------------------------------------ statistics
def test_stats_sample_vs_garmin():
    st = stats([768.1, 764.0, 763.9, 767.4, 762.1, 762.4, 765.6], bullet_gr=139)
    assert st["n"] == 7 and st["avg"] == 764.79
    assert st["sd"] == 2.34 and st["sd_garmin"] == 2.16 and st["es"] == 6.0
    assert st["ci_low"] < st["avg"] < st["ci_high"]
    # mean of per-shot energies, not ½mv² at the mean velocity
    v = [768.1, 764.0, 763.9, 767.4, 762.1, 762.4, 765.6]
    assert st["ke_j"] == round(sum(0.5 * 139 * 0.06479891e-3 * x * x for x in v) / 7)
    assert stats([800, 820], bullet_gr=150)["ke_j"] == round(0.5 * 150 * 0.06479891e-3 * (800**2 + 820**2) / 2)
    assert stats([800, 820])["ke_j"] is None                # no bullet weight → no energy
    assert stats([800, 801])["sd"] is None                 # SD needs n ≥ 3
    assert stats([800, 801, 802, 803])["ci_low"] is None   # CI needs n ≥ 5


def test_temp_sensitivity():
    assert temp_sensitivity([(-12, 326.8), (4.5, 327.4), (16, 334.2), (27.8, 337.9)]) > 0.2
    assert temp_sensitivity([(10, 330), (12, 331), (14, 332)]) is None   # < 10 °C span


# ------------------------------------------------------------------ API end to end
@pytest.fixture()
def client():
    from flask import Flask
    from werkzeug.security import generate_password_hash
    app = Flask(__name__)
    diary.init_app(app)
    conn = diary.connect()
    conn.executescript("DELETE FROM chrono_shots; DELETE FROM chrono_strings; DELETE FROM loads; DELETE FROM guns;")
    diary.kv_set("password_hash", generate_password_hash("test-password-1"), conn)
    diary.upsert_gun(conn, "g-t3x", {"name": "T3X Super Varmint", "cal": "6.5 CM"})
    conn.commit()
    conn.close()
    c = app.test_client()
    assert c.post("/diary/login", data={"password": "test-password-1"}).status_code == 302
    return c


def _upload(c, name, data, path=None):
    return c.post("/diary/api/chrono/parse", headers={"X-Requested-With": "fetch"},
                  data={"files": [(io.BytesIO(data), name)], "paths": [path or name]},
                  content_type="multipart/form-data")


def test_api_parse_commit_idempotent(client):
    data = log_workbook({"13072024": {"date": dt.datetime(2024, 7, 13),
                                      "shots": [("15:33:32", 847.0, 23, 46, 1004), ("15:34:39", 842.6, 23, 46, 1004),
                                                ("15:34:48", 840.7, 23, 46, 1004)]}})
    name = "T3X Super Varmint, 9 g Scenar, Lapua brass, 44.0 gr.xlsx"
    r = _upload(client, name, data, "N555/T3X Super Varmint/" + name)
    assert r.status_code == 200, r.data
    p = r.get_json()
    s = p["strings"][0]
    assert s["suggest"]["gun_id"] == "g-t3x"
    assert s["suggest"]["load"]["name"] == "6.5 CM · Lapua Scenar 139 gr · N555 44.0 gr · Lapua brass"
    item = dict(s, gun_id=s["suggest"]["gun_id"], load=s["suggest"]["load"])
    r = client.post("/diary/api/chrono/commit", json={"strings": [item]})
    assert r.status_code == 200 and len(r.get_json()["created"]) == 1
    again = client.post("/diary/api/chrono/commit", json={"strings": [item]}).get_json()
    assert not again["created"] and again["skipped"][0]["why"] == "already_imported"
    days = client.get("/diary/api/chrono/days?gun=g-t3x").get_json()
    assert days[0]["stats"]["n"] == 3 and days[0]["temp_c"] == 23.0
    lad = client.get("/diary/api/chrono/ladder?gun=g-t3x").get_json()
    assert lad[0]["charge_gr"] == 44.0 and lad[0]["n"] == 3

    # Excluding a shot changes the figures on read
    sid = client.get("/diary/api/chrono/strings").get_json()[0]["id"]
    r = client.patch(f"/diary/api/chrono/strings/{sid}", json={"shots": [{"n": 1, "excluded": True}]})
    assert r.get_json()["stats"]["n"] == 2


def test_api_requires_fetch_header_and_login(client):
    r = client.post("/diary/api/chrono/parse", data={"files": [(io.BytesIO(b"x"), "a.xls")]},
                    content_type="multipart/form-data")
    assert r.status_code == 403
    client.post("/diary/logout")
    assert client.get("/diary/api/chrono/strings").status_code == 401


def test_raw_export_overlap_adds_nothing(client):
    """A raw ShotView session that shares shots with a stored (curated) day must not re-add fliers."""
    conn = diary.connect()
    chrono.commit_strings(conn, [{"date": "2026-10-15", "source": "workbook", "gun_id": "g-t3x",
                                  "shots": [{"v_ms": 836.5, "t": "13:20:00"}, {"v_ms": 836.6, "t": "13:21:00"}]}])
    conn.commit()
    raw = parse_shotview_sheet("x", shotview_rows([836.5, 836.6, 700.0]))
    res = chrono.commit_strings(conn, [dict(raw, gun_id="g-t3x")])
    conn.commit()
    assert not res["created"] and res["skipped"][0]["dropped"] == 1
    assert conn.execute("SELECT count(*) FROM chrono_shots").fetchone()[0] == 2
    conn.close()


def test_gun_matching_by_model_token():
    conn = diary.connect()
    conn.execute("DELETE FROM guns")
    for gid, name in (("g-1", "Tikka T3x Super Varmint"), ("g-2", "Tikka T1x"), ("g-3", "Sako M39"),
                      ("g-4", "WWSD AR-15")):
        diary.upsert_gun(conn, gid, {"name": name})
    conn.commit()
    r = lambda *c: chrono.resolve_gun(conn, list(c))
    assert r("T3X Super Varmint") == "g-1"
    assert r("T1X, SK Long Range, 40,0 gr") == "g-2"           # ShotView title
    assert r("M39 Kisa") == r("M-39 Kisa") == r("M_39") == "g-3"
    assert r("Rifle Bullet 120,0 gr") is None                  # no model token → no guess
    assert r("KP-15") is None                                  # needs an alias
    conn.close()


# ------------------------------------------------------------------ structured loads
from diary import chrono_loads


def test_load_names_and_legacy_migration():
    assert chrono_loads.load_name({"kind": "handload", "bullet_maker": "Lapua", "bullet_model": "Scenar",
                                   "bullet_gr": 139, "brass": "Lapua", "powder": "N555", "charge_gr": 44}) \
        == "Lapua Scenar 139 gr · N555 44.0 gr · Lapua brass"
    assert chrono_loads.load_name({"kind": "factory", "maker": "SK", "product": "Long Range Match", "bullet_gr": 40}) \
        == "SK Long Range Match 40 gr"
    leg = lambda **kw: chrono_loads.from_legacy(kw)["name"]
    assert leg(kind="handload", bullet="9 g Scenar", bullet_gr=139, brass="Lapua", powder="N555", charge_gr=44.0) \
        == "Lapua Scenar 139 gr · N555 44.0 gr · Lapua brass"
    assert leg(kind="handload", bullet="8,8 g Scenar-L", bullet_gr=136, brass="Lapua", powder="N555", charge_gr=44.0) \
        .startswith("Lapua Scenar-L 136 gr")
    assert leg(kind="factory", name="Geco .22lr Rifle 2,6g") == "Geco .22lr Rifle 40 gr"
    assert leg(kind="factory", name="Sellier & Bellot 223. Rem 55gr. FMJ", bullet_gr=123) \
        == "Sellier & Bellot 223. Rem FMJ 55 gr"                     # the name's weight beats a copied cell
    assert leg(kind="factory", name="Sako Range 6,5 9,3g_144g Creedmoor") == "Sako Range 6,5 Creedmoor 144 gr"


def test_migrate_old_loads_table():
    import sqlite3
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.execute("CREATE TABLE loads (id TEXT PRIMARY KEY, name TEXT NOT NULL, kind TEXT, cal TEXT, bullet TEXT, "
                 "bullet_gr REAL, brass TEXT, powder TEXT, charge_gr REAL, notes TEXT, archived INTEGER)")
    conn.execute("INSERT INTO loads VALUES ('l-1','9 g Scenar, Lapua brass, N555 40 gr','handload',NULL,"
                 "'9 g Scenar',139,'Lapua','N555',40,'',0)")
    assert chrono_loads.migrate(conn) == 1
    r = conn.execute("SELECT * FROM loads").fetchone()
    assert r["name"] == "Lapua Scenar 139 gr · N555 40.0 gr · Lapua brass" and r["bullet_model"] == "Scenar"
    assert chrono_loads.migrate(conn) == 0                           # idempotent


def test_api_load_duplicate_and_merge(client):
    base = {"kind": "handload", "bullet_maker": "Lapua", "bullet_model": "Scenar", "bullet_gr": "139",
            "brass": "Lapua", "powder": "N555", "charge_gr": "44,0"}
    r = client.put("/diary/api/chrono/loads/l-a", json=base)
    assert r.status_code == 200 and r.get_json()["name"] == "Lapua Scenar 139 gr · N555 44.0 gr · Lapua brass"
    client.put("/diary/api/chrono/loads/l-b", json=dict(base, charge_gr="43,5"))
    conn = diary.connect()
    chrono.commit_strings(conn, [{"date": "2024-07-13", "source": "workbook", "gun_id": "g-t3x", "load_id": "l-b",
                                  "shots": [{"v_ms": 830.0, "t": "15:00:00"}, {"v_ms": 831.0, "t": "15:01:00"}]}])
    conn.commit()
    conn.close()
    r = client.put("/diary/api/chrono/loads/l-b", json=base)       # edited to the same fields as l-a
    assert r.status_code == 409 and r.get_json()["other"] == "l-a"
    r = client.post("/diary/api/chrono/loads/l-b/merge", json={"into": "l-a"})
    assert r.get_json()["moved"] == 1
    loads = {l["id"]: l for l in client.get("/diary/api/chrono/loads").get_json()}
    assert "l-b" not in loads and loads["l-a"]["use"]["shots"] == 2
    f = client.put("/diary/api/chrono/loads/l-f", json={"kind": "factory", "maker": "SK", "product": "Long Range Match",
                                                         "bullet_gr": "40", "brass": "ignored"}).get_json()
    assert f["name"] == "SK Long Range Match 40 gr" and f["brass"] is None


def test_calibre_is_part_of_the_load(client):
    fmj = {"kind": "handload", "bullet_maker": "Lapua", "bullet_model": "FMJ", "bullet_gr": "123", "brass": "Lapua",
           "powder": "N133", "charge_gr": "43"}
    a = client.put("/diary/api/chrono/loads/l-308", json=dict(fmj, cal=".308 Win"))
    b = client.put("/diary/api/chrono/loads/l-53r", json=dict(fmj, cal="7.62x53R"))
    assert a.status_code == b.status_code == 200                      # not duplicates
    assert a.get_json()["name"] == ".308 Win · Lapua FMJ 123 gr · N133 43.0 gr · Lapua brass"
    assert b.get_json()["name"].startswith("7.62x53R · ")
    assert client.put("/diary/api/chrono/loads/l-x", json=dict(fmj, cal=".308 Win")).status_code == 409


def test_calibre_backfilled_from_gun():
    conn = diary.connect()
    conn.executescript("DELETE FROM chrono_shots; DELETE FROM chrono_strings; DELETE FROM loads;")
    diary.upsert_gun(conn, "g-m39", {"name": "Sako M39", "cal": "7.62x53R"})
    conn.execute("INSERT INTO loads(id,name,kind,bullet_model,bullet_gr,brass,powder,charge_gr) "
                 "VALUES('l-1','x','handload','FMJ',123,'Lapua','N133',43)")
    conn.execute("INSERT INTO chrono_strings(id,date,gun_id,load_id,source,created_at) "
                 "VALUES('c-1','2025-01-09','g-m39','l-1','workbook','now')")
    assert chrono_loads.backfill_calibres(conn) == 1
    r = conn.execute("SELECT cal, name FROM loads WHERE id='l-1'").fetchone()
    assert r["cal"] == "7.62x53R" and r["name"] == "7.62x53R · FMJ 123 gr · N133 43.0 gr · Lapua brass"
    conn.close()


def test_names_regenerated_in_new_order():
    """Loads saved with the old order (… · brass · powder) are renamed on start."""
    import sqlite3
    from diary import chrono_loads
    conn = sqlite3.connect(":memory:")
    conn.executescript("CREATE TABLE guns(id TEXT PRIMARY KEY, data TEXT); CREATE TABLE chrono_strings(id TEXT, load_id TEXT, gun_id TEXT);"
                       "CREATE TABLE loads(id TEXT PRIMARY KEY)")
    chrono_loads.migrate(conn)
    conn.execute("INSERT INTO loads(id,name,kind,cal,bullet_maker,bullet_model,bullet_gr,brass,powder,charge_gr) VALUES"
                 "('l-1','6.5 CM · Lapua Scenar 139 gr · Lapua brass · N555 44.0 gr','handload','6.5 CM','Lapua','Scenar',139,'Lapua','N555',44)")
    assert chrono_loads.backfill_calibres(conn) == 1
    assert conn.execute("SELECT name FROM loads").fetchone()[0] == "6.5 CM · Lapua Scenar 139 gr · N555 44.0 gr · Lapua brass"
